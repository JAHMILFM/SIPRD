"""Flujos funcionales implementados. PostgreSQL con rollback y adjuntos temporales."""
import base64
import io
import tempfile
from pathlib import Path
import unittest
from datetime import date
from uuid import uuid4
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from backend.app.main import app
from backend.app.core.db import engine, get_db
from backend.app.core.config import settings
from backend.app.models import (Usuario, Cliente, Pedido, Ruta, Parada, ReglaCliente,
    Planificacion, PlanificacionVersion, Cobro)

PNG=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZIcAAAAASUVORK5CYII=')

class RequisitosFuncionales(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app);self.client.__enter__()
        async def abrir():
            conn=await engine.connect();tx=await conn.begin()
            db=AsyncSession(bind=conn,expire_on_commit=False,join_transaction_mode='create_savepoint')
            return conn,tx,db
        self.conn,self.tx,self.db=self.client.portal.call(abrir)
        async def database():yield self.db
        app.dependency_overrides[get_db]=database
        self.temp=tempfile.TemporaryDirectory(prefix="rf-",dir=Path(".local").resolve());self.upload=settings.UPLOAD_DIR;settings.UPLOAD_DIR=self.temp.name
        self.headers={}
        for name,pwd in [('dhuerta','jefe123'),('lpomalaya','dist123'),('admin.ti','ti2026'),('elopez','rep123'),('tesoreria','teso123')]:
            self.headers[name]=self.login(name,pwd)
        self.tag=uuid4().hex[:10]

    def tearDown(self):
        async def cerrar():
            await self.db.close();await self.tx.rollback();await self.conn.close()
        self.client.portal.call(cerrar);app.dependency_overrides.clear()
        settings.UPLOAD_DIR=self.upload
        self.client.__exit__(None,None,None)
        if not Path(self.temp.name).resolve().is_relative_to(Path('.local').resolve()):
            raise RuntimeError('Directorio temporal fuera del proyecto')
        self.temp.cleanup()

    def call(self,method,path,user='dhuerta',expected=200,**kwargs):
        r=self.client.request(method,'/api/v1'+path,headers=self.headers[user],**kwargs)
        self.assertEqual(r.status_code,expected,r.text)
        return r

    def login(self,name,pwd):
        r=self.client.post('/api/v1/auth/login',json={'usuario':name,'clave':pwd})
        self.assertEqual(r.status_code,200,r.text)
        return {'Authorization':'Bearer '+r.json()['access_token']}

    def nuevo_usuario(self,rol='REPARTIDOR'):
        data=dict(nombre='Prueba',apellidos='Funcional',documento='ID-'+self.tag,correo=f'correo-{self.tag}@example.com',usuario='usuario-'+self.tag,clave='Prueba123',rol=rol)
        r=self.call('POST','/usuarios','admin.ti',201,json=data)
        return r.json()['id'],data

    async def fixture_pedidos(self,plan=False):
        c=Cliente(codigo_externo='QA-'+self.tag,razon_social='Cliente QA',direccion='Lurín',distrito='Lurín',latitud=-12.275,longitud=-76.872)
        self.db.add(c);await self.db.flush()
        p=Pedido(codigo_externo='QA-P-'+self.tag,id_cliente=c.id_cliente,fecha_programada=date(2026,10,12),peso_kg=100,volumen_m3=1,importe=250,estado='PENDIENTE')
        self.db.add(p);await self.db.flush()
        if not plan:return c.id_cliente,p.id_pedido
        pl=Planificacion(fecha='12/10/2026',version_vigente=1,creada_por='1');self.db.add(pl);await self.db.flush()
        v=PlanificacionVersion(planificacion_id=pl.id,numero=1,estado='CONFIRMADA');self.db.add(v);await self.db.flush()
        ruta=Ruta(version_id=v.id,fecha_ruta=date(2026,10,12),id_vehiculo=1,id_repartidor=4,estado='PUBLICADA');self.db.add(ruta);await self.db.flush()
        parada=Parada(id_ruta=ruta.id_ruta,id_pedido=p.id_pedido,secuencia=1,estado='PENDIENTE');self.db.add(parada);p.estado='PLANIFICADO';await self.db.commit()
        return c.id_cliente,p.id_pedido,parada.id_parada

    def test_01_acceso_y_ciclo_administracion(self):
        """RF-ACC-01; RF-ADM-01/02/03/04; revocación y cambio de rol efectivos."""
        self.assertEqual(self.client.get('/api/v1/usuarios').status_code,401)
        self.assertEqual(self.client.post('/api/v1/auth/login',json={'usuario':'dhuerta','clave':'mal'}).status_code,401)
        id,data=self.nuevo_usuario('ASISTENTE')
        self.headers['nuevo']=self.login(data['usuario'],data['clave'])
        self.call('GET','/planificaciones','nuevo')
        rows=self.call('GET','/usuarios','admin.ti').json();u=next(x for x in rows if x['id']==id)
        self.assertEqual(u['documento'],data['documento']);self.assertEqual(u['usuario'],data['usuario'])
        self.assertNotIn('password_hash',u)
        self.call('POST','/usuarios','dhuerta',403,json={**data,'rol':'ADMINISTRADOR'})
        self.call('POST','/usuarios','admin.ti',409,json=data)
        data.update(nombre='Nombre actualizado',rol='TESORERIA',activo=True,clave=None)
        self.call('PUT',f'/usuarios/{id}','admin.ti',json=data)
        self.call('GET','/planificaciones','nuevo',403)
        self.call('GET','/cobros','nuevo')
        self.call('POST',f'/usuarios/{id}/desactivar','admin.ti')
        self.call('GET','/cobros','nuevo',401)
        self.assertEqual(self.client.post('/api/v1/auth/login',json={'usuario':data['usuario'],'clave':'Prueba123'}).status_code,403)
        self.call('POST',f'/usuarios/{id}/reactivar','admin.ti')
        self.call('GET','/cobros','nuevo')

    def test_02_vehiculos_y_reglas_persistentes(self):
        """RF-VEH-01/02/03/04; RF-RES-01/02/03/04."""
        v=dict(placa='QA-'+self.tag,capacidad_peso_kg=2000,capacidad_volumen_m3=8,conductor='Chofer QA',estado='MANTENIMIENTO')
        id=self.call('POST','/vehiculos','lpomalaya',201,json=v).json()['id']
        disponibles=self.call('GET','/vehiculos?disponibles=true').json()
        self.assertNotIn(id,[x['id'] for x in disponibles])
        self.call('PUT',f'/datos/vehiculos/{id}','lpomalaya',json=dict(pesoMax=3.5,volMax=9,conductor='Chofer actualizado',estado='DISPONIBLE'))
        row=next(x for x in self.call('GET','/vehiculos').json() if x['id']==id)
        self.assertEqual(row['conductor'],'Chofer actualizado');self.assertEqual(row['capacidad_kg'],3500)
        self.call('POST',f'/vehiculos/{id}/desactivar','lpomalaya')
        self.assertNotIn(id,[x['id'] for x in self.call('GET','/vehiculos?disponibles=true').json()])
        self.call('POST','/vehiculos','elopez',403,json=v)
        cid,pid=self.client.portal.call(self.fixture_pedidos)
        req=dict(cliente_id=cid,dias=[1,2,3],ventana='09:00–15:00')
        rid=self.call('POST','/datos/reglas','lpomalaya',201,json=req).json()['id']
        rows=self.call('GET',f'/datos/reglas?cliente_id={cid}').json();self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['dias'],[1,2,3]);self.assertEqual(rows[0]['ventana'],'09:00–15:00')
        req.update(dias=[1,6],ventana='10:00–14:00')
        self.call('PUT',f'/datos/reglas/{rid}','lpomalaya',json=req)
        self.call('PUT',f'/datos/reglas/{rid}','lpomalaya',422,json={**req,'ventana':'17:00–09:00'})
        self.call('POST',f'/datos/reglas/{rid}/desactivar','lpomalaya')
        self.assertEqual(self.call('GET',f'/datos/reglas?cliente_id={cid}').json(),[])
        rows=self.call('GET',f'/datos/reglas?cliente_id={cid}&incluir_inactivas=true').json()
        self.assertEqual(rows[0]['activa'],False)

    def test_03_propuesta_restricciones_y_confirmacion(self):
        """RF-PLAN-01/02/04; día, capacidad, asignación, confirmación y ejecución."""
        cid,pid=self.client.portal.call(self.fixture_pedidos)
        async def adicionales():
            c=Cliente(codigo_externo='QA-B-'+self.tag,razon_social='Cliente cerrado',direccion='Lurín',latitud=-12.275,longitud=-76.872);self.db.add(c);await self.db.flush()
            self.db.add(ReglaCliente(id_cliente=c.id_cliente,tipo_regla='DIAS_DISPONIBLES',valor='2',activa=True))
            self.db.add(Pedido(codigo_externo='QA-BP-'+self.tag,id_cliente=c.id_cliente,fecha_programada=date(2026,10,12),peso_kg=10,volumen_m3=1,importe=100))
            self.db.add(Pedido(codigo_externo='QA-PESADO-'+self.tag,id_cliente=cid,fecha_programada=date(2026,10,12),peso_kg=999999,volumen_m3=1,importe=100));await self.db.commit()
        self.client.portal.call(adicionales)
        r=self.call('POST','/planificaciones','lpomalaya',201,json={'fecha':'12/10/2026','vehiculo_ids':[1]}).json()
        plan=r['planificacion_id'];base=f'/planificaciones/{plan}/versiones/1'
        detalle=self.call('GET',base,'lpomalaya').json()
        self.assertEqual(len(detalle['rutas']),1);self.assertEqual(len(detalle['no_asignados']),2)
        ruta=detalle['rutas'][0];self.assertIn('restricciones',ruta['paradas'][0])
        self.call('POST',base+'/confirmar','lpomalaya',409)
        self.call('PUT',base+f'/rutas/{ruta["id"]}/repartidor','lpomalaya',422,json={'repartidor_id':5})
        self.call('PUT',base+f'/rutas/{ruta["id"]}/repartidor','lpomalaya',json={'repartidor_id':4})
        self.call('POST',base+'/confirmar','elopez',403)
        self.call('PUT','/datos/vehiculos/1',json=dict(pesoMax=.01,volMax=32,conductor='Chofer',estado='DISPONIBLE'))
        self.call('POST',base+'/confirmar','lpomalaya',409)
        self.call('PUT','/datos/vehiculos/1',json=dict(pesoMax=12,volMax=32,conductor='Chofer',estado='DISPONIBLE'))
        self.call('POST',base+'/confirmar','lpomalaya')
        self.assertEqual(self.call('GET',base).json()['estado'],'CONFIRMADA')
        rutas=self.call('GET','/reparto/mi-ruta','elopez').json()['rutas']
        self.assertIn(str(ruta['id']),[x['id'] for x in rutas])
        self.call('PUT',base+f'/rutas/{ruta["id"]}/repartidor','lpomalaya',409,json={'repartidor_id':4})

    def test_04_entrega_incidencia_cobro_contraste_y_excel(self):
        """RF-REP-01/02; RF-INC-01/02/03; RF-COB-01/02/03/04/05/07."""
        cid,pid,par=self.client.portal.call(self.fixture_pedidos,True)
        id,data=self.nuevo_usuario();self.headers['ajeno']=self.login(data['usuario'],data['clave'])
        self.call('GET',f'/reparto/paradas/{par}/evidencias','ajeno',403)
        self.call('POST',f'/reparto/paradas/{par}/evidencias','ajeno',403,files={'archivo':('foto.png',PNG,'image/png')})
        self.call('POST',f'/reparto/paradas/{par}/evidencias','elopez',400,files={'archivo':('foto.png',b'no es una imagen','image/png')})
        self.call('POST','/cobros','elopez',409,json={'pedido_id':str(pid),'importe':250,'medio_pago':'EFECTIVO'})
        ev=self.call('POST',f'/reparto/paradas/{par}/evidencias','elopez',201,files={'archivo':('foto.png',PNG,'image/png')}).json()['evidencia_id']
        self.assertTrue(self.call('GET',f'/reparto/paradas/{par}/evidencias/{ev}/archivo','elopez').content.startswith(b'\x89PNG'))
        self.call('POST',f'/reparto/paradas/{par}/incidencias','elopez',201,data={'tipo':'OTRO','descripcion':'Demora por tráfico'})
        inc=next(x for x in self.call('GET','/incidencias','lpomalaya').json() if x['parada_id']==str(par))
        self.call('PATCH',f'/incidencias/{inc["id"]}','lpomalaya',422,json={'estado':'RESUELTA'})
        self.call('PATCH',f'/incidencias/{inc["id"]}','lpomalaya',json={'estado':'RESUELTA','resolucion':'Entrega coordinada'})
        actualizado=next(x for x in self.call('GET','/incidencias').json() if x['id']==inc['id']);self.assertEqual(actualizado['resolucion'],'Entrega coordinada')
        self.call('POST','/cobros','ajeno',403,json={'pedido_id':str(pid),'importe':250,'medio_pago':'EFECTIVO'})
        self.call('POST','/cobros','elopez',409,json={'pedido_id':str(pid),'importe':250,'medio_pago':'YAPE'})
        observado=self.call('POST','/cobros','elopez',201,json={'pedido_id':str(pid),'importe':20,'medio_pago':'YAPE','numero_operacion':'QA-'+self.tag}).json()['cobro_id']
        cobro=self.call('POST','/cobros','elopez',201,json={'pedido_id':str(pid),'importe':250,'medio_pago':'EFECTIVO'}).json()['cobro_id']
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',409,json={'resultado':'CONFORME','movimiento_bancario':'MOV-QA'})
        for id in [observado,cobro]:self.call('POST',f'/cobros/{id}/comprobante','elopez',201,files={'archivo':('comprobante.png',PNG,'image/png')})
        self.call('GET',f'/cobros/{cobro}/comprobante/archivo','ajeno',403)
        self.call('GET',f'/cobros/{cobro}/comprobante/archivo','tesoreria')
        self.call('POST',f'/cobros/{observado}/contraste','tesoreria',409,json={'resultado':'OBSERVADO','movimiento_bancario':'MOV-QA'})
        self.call('POST',f'/cobros/{observado}/contraste','tesoreria',json={'resultado':'OBSERVADO','movimiento_bancario':'MOV-QA','observacion':'No coincide con el movimiento bancario'})
        rows=self.call('GET','/cobros?estado=OBSERVADO','tesoreria').json()
        self.assertEqual(next(x for x in rows if x['id']==observado)['observaciones'][0]['texto'],'No coincide con el movimiento bancario')
        self.call('POST',f'/cobros/{cobro}/contraste','dhuerta',403,json={'resultado':'CONFORME','movimiento_bancario':'MOV-QA'})
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',json={'resultado':'CONFORME','movimiento_bancario':'ARQUEO-QA'})
        async def estado():
            self.db.expire_all();return (await self.db.get(Pedido,pid)).estado
        self.assertEqual(self.client.portal.call(estado),'CERRADO')
        response=self.call('GET','/cobros/exportacion.xlsx','tesoreria')
        wb=openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Detalle de Cobros',wb.sheetnames)
        self.assertTrue(any(row[1]=='QA-P-'+self.tag and row[7]=='CONFORME' for row in wb['Detalle de Cobros'].iter_rows(min_row=2,values_only=True)))
        self.call('GET','/cobros/exportacion.xlsx','elopez',403)
        audit=self.call('GET',f'/auditoria?entidad_id={cobro}','admin.ti').json()
        self.assertTrue(any(a['entidad_id']==cobro and a['valores_despues'].get('movimiento_bancario')=='ARQUEO-QA' for a in audit))

    def test_05_cierre_exige_monto_exacto(self):
        """Un cobro superior al importe no debe cerrar el pedido."""
        cid,pid,par=self.client.portal.call(self.fixture_pedidos,True)
        self.call('POST',f'/reparto/paradas/{par}/evidencias','elopez',201,files={'archivo':('foto.png',PNG,'image/png')})
        cobro=self.call('POST','/cobros','elopez',201,json={'pedido_id':str(pid),'importe':251,'medio_pago':'EFECTIVO'}).json()['cobro_id']
        self.call('POST',f'/cobros/{cobro}/comprobante','elopez',201,files={'archivo':('pago.png',PNG,'image/png')})
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',json={'resultado':'CONFORME','movimiento_bancario':'ARQUEO-QA'})
        async def estado():
            self.db.expire_all();return (await self.db.get(Pedido,pid)).estado
        self.assertEqual(self.client.portal.call(estado),'ENTREGADO')
        self.call('POST',f'/pedidos/{pid}/cerrar','tesoreria',409)

    def test_06_incidencias_auditoria_y_nombres(self):
        """Consulta general y filtrada sin lazy loading; identidad real e idempotencia."""
        incidencias = self.call('GET', '/incidencias').json()
        self.assertGreaterEqual(len(incidencias), 3)
        primera = incidencias[0]
        filtradas = self.call('GET', f"/incidencias?ruta_id={primera['ruta_id']}").json()
        self.assertTrue(all(i['ruta_id'] == primera['ruta_id'] for i in filtradas))
        self.call('GET', '/incidencias?ruta_id=no-es-un-id', expected=422)
        self.call('GET', '/incidencias', 'elopez', 403)
        self.call('PATCH', f"/incidencias/{primera['id']}", json={'estado':'EN_REVISION','resolucion':'Revisión de prueba'})
        audit = self.call('GET', f"/auditoria?entidad=incidencia&entidad_id={primera['id']}").json()
        self.assertEqual(audit[0]['usuario_nombre'], 'Dennys Huerta')
        self.assertEqual(audit[0]['usuario_login'], 'dhuerta')
        self.assertEqual(audit[0]['usuario_rol'], 'JEFE')
        self.assertEqual(audit[0]['usuario_correo'], 'dhuerta@alfadistribuidores.com')
        usuarios = self.call('GET', '/usuarios', 'admin.ti').json()
        self.assertTrue(all(u['nombre'] and u['usuario'] and 'activo' in u for u in usuarios))
        async def comprobar_nombres():
            from sqlalchemy import func
            from backend.app.datos_demo import actualizar_nombres_clientes, NOMBRES_CLIENTES
            antes = await self.db.scalar(select(func.count()).select_from(Usuario))
            cliente = await self.db.scalar(select(Cliente).where(Cliente.codigo_externo == 'DEMO-CLI-001'))
            cliente.razon_social = 'Cliente simulado 01'
            await self.db.flush()
            cambios = await actualizar_nombres_clientes(self.db)
            self.assertEqual(len(cambios), 1)
            self.assertEqual(cliente.razon_social, NOMBRES_CLIENTES[0])
            self.assertEqual(await actualizar_nombres_clientes(self.db), [])
            cliente.razon_social = 'Nombre personalizado por el usuario'
            await self.db.flush()
            self.assertEqual(await actualizar_nombres_clientes(self.db), [])
            self.assertEqual(cliente.razon_social, 'Nombre personalizado por el usuario')
            self.assertEqual(antes, await self.db.scalar(select(func.count()).select_from(Usuario)))
        self.client.portal.call(comprobar_nombres)

    def test_07_jefe_cuentas_operativas_y_consulta_financiera(self):
        rows=self.call('GET','/usuarios').json()
        administrador=next(u for u in rows if u['rol']=='ADMINISTRADOR')
        self.call('PUT',f"/usuarios/{administrador['id']}",expected=403,json={**administrador,'clave':None})
        self.call('POST',f"/usuarios/{administrador['id']}/desactivar",expected=403)
        id,data=self.nuevo_usuario('ASISTENTE')
        data.update(nombre='Operativo editado',activo=True,clave=None)
        self.call('PUT',f'/usuarios/{id}',json=data)
        self.call('POST',f'/usuarios/{id}/desactivar')
        self.call('POST',f'/usuarios/{id}/reactivar')
        self.assertEqual(next(u for u in self.call('GET','/usuarios').json() if u['id']==id)['nombre'],'Operativo editado')
        self.assertGreater(len(self.call('GET','/cobros').json()),0)
        self.call('GET','/cobros/exportacion.xlsx')
        self.call('GET','/usuarios','elopez',403)

    def test_08_edicion_manual_y_versiones_sin_perder_historial(self):
        cid,pid=self.client.portal.call(self.fixture_pedidos)
        async def otro_pedido():
            segundo=Pedido(codigo_externo='QA-SEG-'+self.tag,id_cliente=cid,fecha_programada=date(2026,10,12),peso_kg=100,volumen_m3=1,importe=250,estado='PENDIENTE')
            self.db.add(segundo);await self.db.flush();await self.db.commit();return segundo.id_pedido
        pid2=self.client.portal.call(otro_pedido)
        r=self.call('POST','/planificaciones',expected=201,json={'fecha':'12/10/2026','vehiculo_ids':[1]}).json()
        plan=r['planificacion_id'];base=f'/planificaciones/{plan}/versiones/1'
        original=self.call('GET',base).json()
        self.call('PUT',base+'/asignaciones',expected=422,json={'rutas':[{'vehiculo_id':1,'pedido_ids':[pid,pid]}]})
        self.assertEqual(self.call('GET',base).json()['rutas'][0]['id'],original['rutas'][0]['id'])
        # Infeasible edits must preserve the previous routes atomically.
        self.call('PUT','/datos/vehiculos/2',json=dict(pesoMax=.01,volMax=32,conductor='Chofer',estado='DISPONIBLE'))
        self.call('PUT',base+'/asignaciones',expected=409,json={'rutas':[{'vehiculo_id':2,'repartidor_id':4,'pedido_ids':[pid2,pid]}]})
        self.assertEqual(self.call('GET',base).json()['rutas'][0]['id'],original['rutas'][0]['id'])
        self.call('PUT','/datos/vehiculos/2',json=dict(pesoMax=12,volMax=32,conductor='Chofer',estado='DISPONIBLE'))
        self.call('PUT',base+'/asignaciones',json={'rutas':[{'vehiculo_id':2,'repartidor_id':4,'pedido_ids':[pid2,pid]}]})
        editado=self.call('GET',base).json()
        self.assertEqual(editado['rutas'][0]['vehiculo']['id'],'2')
        self.assertEqual([p['pedido_id'] for p in editado['rutas'][0]['paradas']],[str(pid2),str(pid)])
        self.assertEqual(editado['rutas'][0]['peso_kg'],200)
        self.call('POST',base+'/confirmar')
        self.call('PUT',base+'/asignaciones',expected=409,json={'rutas':[{'vehiculo_id':1,'pedido_ids':[pid]}]})
        nuevo=self.call('POST',base+'/reoptimizar',expected=201,json={'vehiculo_ids':[1],'motivo':'Cambiar vehículo por mantenimiento'}).json()
        self.assertEqual(nuevo['version_numero'],2)
        base2=f'/planificaciones/{plan}/versiones/2'
        historial=self.call('GET',f'/planificaciones/{plan}/versiones').json()
        self.assertEqual([v['numero'] for v in historial],[2,1])
        self.assertEqual(self.call('GET',base).json()['estado'],'CONFIRMADA')
        ruta2=self.call('GET',base2).json()['rutas'][0]
        self.call('PUT',base2+f"/rutas/{ruta2['id']}/repartidor",json={'repartidor_id':4})
        self.call('POST',base2+'/confirmar')
        self.assertEqual(self.call('GET',base).json()['estado'],'SUPERADA')
        self.assertEqual([p['pedido_id'] for p in self.call('GET',base).json()['rutas'][0]['paradas']],[str(pid2),str(pid)])
        self.call('POST',base+'/confirmar',expected=409)
        self.call('POST',f"/reparto/paradas/{ruta2['paradas'][0]['id']}/evidencias",'elopez',201,files={'archivo':('foto.png',PNG,'image/png')})
        self.call('POST',base2+'/reoptimizar',expected=409,json={'vehiculo_ids':[1],'motivo':'Jornada iniciada'})

    def test_09_subsanacion_comprobante_y_recontraste(self):
        cid,pid,par=self.client.portal.call(self.fixture_pedidos,True)
        self.call('POST',f'/reparto/paradas/{par}/evidencias','elopez',201,files={'archivo':('foto.png',PNG,'image/png')})
        cobro=self.call('POST','/cobros','elopez',201,json={'pedido_id':str(pid),'importe':20,'medio_pago':'YAPE','numero_operacion':'ERR-'+self.tag}).json()['cobro_id']
        self.call('POST',f'/cobros/{cobro}/comprobante','elopez',201,files={'archivo':('pago.png',PNG,'image/png')})
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',json={'resultado':'OBSERVADO','movimiento_bancario':'MOV-QA','observacion':'Corregir el importe y comprobante'})
        req={'importe':250,'medio_pago':'YAPE','numero_operacion':'OK-'+self.tag,'respuesta':'Importe corregido y comprobante actualizado'}
        self.call('PUT',f'/cobros/{cobro}/correccion','tesoreria',403,json=req)
        id,data=self.nuevo_usuario();self.headers['ajeno']=self.login(data['usuario'],data['clave'])
        self.call('PUT',f'/cobros/{cobro}/correccion','ajeno',403,json=req)
        self.call('PUT',f'/cobros/{cobro}/correccion','elopez',json=req)
        c=next(x for x in self.call('GET','/cobros','tesoreria').json() if x['id']==cobro)
        self.assertEqual(c['estado'],'PENDIENTE_CONTRASTE');self.assertFalse(c['tiene_comprobante'])
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',409,json={'resultado':'CONFORME','movimiento_bancario':'MOV-QA'})
        self.call('POST',f'/cobros/{cobro}/comprobante','elopez',201,files={'archivo':('corregido.png',PNG,'image/png')})
        self.call('POST',f'/cobros/{cobro}/contraste','tesoreria',json={'resultado':'CONFORME','movimiento_bancario':'MOV-QA-CORREGIDO'})
        c=next(x for x in self.call('GET','/cobros','tesoreria').json() if x['id']==cobro)
        self.assertEqual(c['comprobante_version'],2);self.assertTrue(c['observaciones'][0]['resuelta'])
        async def estado():
            self.db.expire_all();return (await self.db.get(Pedido,pid)).estado
        self.assertEqual(self.client.portal.call(estado),'CERRADO')
        self.call('PUT',f'/cobros/{cobro}/correccion','elopez',409,json=req)

    def test_10_cobranzas_jefe_relaciones_horarios_y_etiquetas(self):
        datos=self.call('GET','/datos').json()
        self.assertGreaterEqual(len(datos['cobros']),9)
        async def verificar():
            from backend.app.datos_demo import demo_id, completar_cobranzas_jornada
            from backend.app.models import ParadaRuta, ObservacionCobro
            c=await self.db.get(Cobro,demo_id('cobro-0'))
            p=await self.db.get(Pedido,c.pedido_id)
            parada=await self.db.get(ParadaRuta,c.parada_id)
            ruta=await self.db.get(Ruta,parada.id_ruta)
            cliente=await self.db.get(Cliente,p.id_cliente)
            self.assertEqual(c.registrado_por,ruta.id_repartidor)
            self.assertEqual(parada.id_pedido,p.id_pedido)
            self.assertEqual(c.monto_cobrado,p.importe)
            self.assertEqual(p.estado,'CERRADO')
            self.assertEqual(cliente.razon_social,'Botica Las Palmeras')
            self.assertGreater(c.fecha_registro,parada.hora_llegada_real)
            self.assertNotIn('simulado',c.nota.lower())
            self.assertIsNone(c.numero_operacion)
            fila=next(x for x in datos['cobros'] if x['id']==c.id)
            self.assertEqual(fila['ruta'],str(ruta.id_ruta))
            self.assertEqual(fila['hora'],'08:15')
            self.assertEqual(await completar_cobranzas_jornada(self.db),[])
            c.nota='Anotación ingresada por el usuario'
            await self.db.flush()
            self.assertEqual(await completar_cobranzas_jornada(self.db),[])
            self.assertEqual(c.nota,'Anotación ingresada por el usuario')
        self.client.portal.call(verificar)

if __name__=='__main__':unittest.main()
